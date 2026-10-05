"""Offline tests. Every API is faked, including the moment it runs out of
credits, so we can prove the fallback works. Run: python -m pytest -q"""
import json
import os
import sys

import pytest

sys.path.insert(0, os.path.dirname(os.path.dirname(__file__)))

from finder import core
from finder.export import to_csv, to_xlsx
from finder.pipeline import (CATCH_ALL, FOUND, GUESS, VERIFIED, Pipeline, Settings, Target)
from finder.providers import Registry
from finder.usage import UsageStore


# ---------------------------------------------------------------- fakes
class Resp:
    def __init__(self, status=200, body=None):
        self.status_code = status
        self._body = body if body is not None else {}
        self.text = json.dumps(self._body)

    def json(self):
        return self._body


class FakeHTTP:
    """Routes requests by URL substring to handler functions."""
    def __init__(self, routes):
        self.routes = routes
        self.calls = []

    def request(self, method, url, **kw):
        self.calls.append(url)
        for frag, handler in self.routes.items():
            if frag in url:
                return handler(kw)
        raise AssertionError(f"unexpected call to {url}")

    def count(self, frag):
        return sum(1 for u in self.calls if frag in u)


KEYS = {k: "test" for k in ["SERPER_API_KEY", "TAVILY_API_KEY", "PROSPEO_API_KEY",
                            "HUNTER_API_KEY", "TOMBA_API_KEY", "TOMBA_SECRET",
                            "REOON_API_KEY", "ZEROBOUNCE_API_KEY"]}

LINKEDIN_ROWS = [
    {"title": "Ananya Rao - Category Manager - Acme | LinkedIn", "link": "https://in.linkedin.com/in/ananya-rao",
     "snippet": "Category Manager at Acme."},
    {"title": "Rohan Gupta - Senior Category Manager - Acme | LinkedIn", "link": "https://in.linkedin.com/in/rohan-g",
     "snippet": "Senior Category Manager at Acme. Ex-Flipkart."},
    {"title": "Meera Iyer - Software Engineer - Acme | LinkedIn", "link": "https://in.linkedin.com/in/meera",
     "snippet": "Engineer at Acme."},
]
REAL_MAILBOXES = {"arao@acme.com", "rgupta@acme.com"}  # Acme uses {f}{last}


def reoon_ok(kw):
    email = kw["params"]["email"]
    return Resp(200, {"status": "safe" if email in REAL_MAILBOXES else "invalid"})


def serper_ok(kw):
    q = kw["json"]["q"]
    if q.startswith("site:linkedin.com/in"):
        return Resp(200, {"organic": LINKEDIN_ROWS})
    return Resp(200, {"organic": []})


def prospeo_ok(kw):
    d = kw["json"]["data"]
    email = f"{d['first_name'][0]}{d['last_name']}@acme.com".lower()
    return Resp(200, {"error": False, "person": {"email": {"email": email, "status": "VERIFIED", "revealed": True}}})


@pytest.fixture
def usage(tmp_path):
    return UsageStore(tmp_path / "usage.json")


def make(routes, usage, mode="saver", smtp=False):
    http = FakeHTTP(routes)
    reg = Registry(KEYS, usage, log=lambda m: None, http=http, use_ddg=False, use_smtp=smtp)
    pipe = Pipeline(reg, Settings(mode=mode, crawl=False, snippets=False))
    return pipe, http


# ---------------------------------------------------------------- core
def test_names_and_patterns():
    assert core.split_name("Dr. Arjun K. Mehta, MBA") == ("arjun", "mehta")
    assert core.pattern_from_email("psharma@acme.com", "Priya Sharma") == "{f}{last}"
    ev = core.PatternEvidence()
    core.learn_from_email(ev, "neha.iyer@acme.com", "Grievance Officer: Neha Iyer", [])
    assert ev.confirmed() == "{first}.{last}"


def test_linkedin_parsing_filters_role_and_company():
    rows = [{"title": r["title"], "url": r["link"], "body": r["snippet"]} for r in LINKEDIN_ROWS]
    people = [core.parse_linkedin_result(r, "Acme", "Category Manager") for r in rows]
    assert [p.name if p else None for p in people] == ["Ananya Rao", "Rohan Gupta", None]


# ---------------------------------------------------------------- pipeline
def test_saver_mode_spends_one_finder_credit_then_learns(usage):
    pipe, http = make({"serper.dev": serper_ok, "prospeo.io": prospeo_ok,
                       "reoon.com": reoon_ok}, usage)
    rows = pipe.run([Target("Acme", "acme.com", "Category Manager", limit=5)])
    assert [r.email for r in rows] == ["arao@acme.com", "rgupta@acme.com"]
    assert all(r.status == VERIFIED for r in rows)
    # first person: Prospeo. Second: pattern learned, confirmed by Reoon. One credit.
    assert http.count("prospeo.io") == 1
    assert rows[1].method == "Pattern + Reoon"
    assert "may no longer work there" not in rows[1].note  # "Ex-Flipkart" isn't ex-Acme
    assert usage.used("prospeo", "month") == 1


def test_finder_falls_back_when_prospeo_runs_out(usage):
    def prospeo_empty(kw):
        return Resp(400, {"error": True, "error_code": "INSUFFICIENT_CREDITS"})

    def hunter(kw):
        p = kw["params"]
        return Resp(200, {"data": {"email": f"{p['first_name'][0]}{p['last_name']}@acme.com",
                                   "score": 92, "verification": {"status": "valid"}}})

    pipe, http = make({"serper.dev": serper_ok, "prospeo.io": prospeo_empty,
                       "hunter.io": hunter, "reoon.com": reoon_ok}, usage)
    rows = pipe.run([Target("Acme", "acme.com", "Category Manager")])
    assert rows[0].email == "arao@acme.com" and rows[0].method == "Hunter"
    assert usage.is_exhausted("prospeo")
    assert any("Prospeo is out of free credits" in e for e in pipe.events)
    # Prospeo stays off for the rest of the month: second run doesn't call it
    pipe2, http2 = make({"serper.dev": serper_ok, "hunter.io": hunter,
                         "reoon.com": reoon_ok}, usage)
    pipe2.run([Target("Acme", "acme.com", "Category Manager")])
    assert http2.count("prospeo.io") == 0


def test_search_falls_back_from_serper_to_tavily(usage):
    def serper_out(kw):
        return Resp(400, {"message": "Not enough credits"})

    def tavily(kw):
        assert kw["json"]["include_domains"] == ["linkedin.com"]
        return Resp(200, {"results": [{"title": r["title"], "url": r["link"], "content": r["snippet"]}
                                      for r in LINKEDIN_ROWS]})

    pipe, http = make({"serper.dev": serper_out, "tavily.com": tavily,
                       "prospeo.io": prospeo_ok, "reoon.com": reoon_ok}, usage)
    rows = pipe.run([Target("Acme", "acme.com", "Category Manager")])
    assert len(rows) == 2
    assert usage.is_exhausted("serper")


def test_verifier_falls_back_reoon_to_zerobounce(usage):
    def reoon_out(kw):
        return Resp(200, {"status": "error", "reason": "You do not have enough credits"})

    def zb(kw):
        return Resp(200, {"status": "valid" if kw["params"]["email"] in REAL_MAILBOXES else "invalid"})

    targets = [Target("Acme", "acme.com", people=[core.Person.from_name("Ananya Rao")])]
    pipe, http = make({"reoon.com": reoon_out, "zerobounce.net": zb,
                       "prospeo.io": lambda kw: Resp(400, {"error_code": "NO_MATCH"}),
                       "hunter.io": lambda kw: Resp(200, {"data": {"email": None}}),
                       "tomba.io": lambda kw: Resp(200, {"data": {"email": None}})}, usage)
    pipe.s.guesses_to_verify = 3
    rows = pipe.run(targets)
    assert rows[0].email == "arao@acme.com" and rows[0].status == VERIFIED
    assert rows[0].method == "Pattern + ZeroBounce"
    assert usage.is_exhausted("reoon")


def test_catch_all_and_no_verifier(usage):
    targets = [Target("Acme", "acme.com", people=[core.Person.from_name("Ananya Rao")])]
    pipe, _ = make({"reoon.com": lambda kw: Resp(200, {"status": "catch_all"}),
                    "prospeo.io": lambda kw: Resp(400, {"error_code": "NO_MATCH"}),
                    "hunter.io": lambda kw: Resp(200, {"data": {}}),
                    "tomba.io": lambda kw: Resp(200, {"data": {}})}, usage)
    row = pipe.run(targets)[0]
    assert row.status == CATCH_ALL and row.email == "ananya.rao@acme.com"


def test_everything_exhausted_still_returns_best_guess(usage):
    for name in ["prospeo", "hunter", "tomba", "tomba_verify", "reoon", "zerobounce"]:
        usage.mark_exhausted(name)
    targets = [Target("Acme", "acme.com", people=[core.Person.from_name("Ananya Rao")])]
    pipe, http = make({}, usage)
    row = pipe.run(targets)[0]
    assert row.status == GUESS and row.email == "ananya.rao@acme.com"
    assert http.calls == []  # nothing called once everything is out


def test_local_cap_stops_before_the_service_does(usage, monkeypatch):
    monkeypatch.setenv("CAP_PROSPEO", "1")
    targets = [Target("Acme", "acme.com", people=[core.Person.from_name("Ananya Rao"),
                                                  core.Person.from_name("Kabir Shah")])]
    pipe, http = make({"prospeo.io": prospeo_ok, "reoon.com": reoon_ok,
                       "hunter.io": lambda kw: Resp(200, {"data": {}}),
                       "tomba.io": lambda kw: Resp(200, {"data": {}})}, usage, mode="accurate")
    pipe.run(targets)
    assert http.count("prospeo.io") == 1


def test_rate_limit_retries_then_moves_on(usage, monkeypatch):
    monkeypatch.setattr("finder.providers.time.sleep", lambda s: None)
    hits = {"n": 0}

    def tavily_busy(kw):
        hits["n"] += 1
        return Resp(429, {"detail": {"error": "slow down"}})

    def serper_out(kw):
        return Resp(400, {"message": "Not enough credits"})

    pipe, http = make({"serper.dev": serper_out, "tavily.com": tavily_busy}, usage)
    pipe.reg.searchers = pipe.reg.searchers[:2]
    pipe.reg.search.providers = pipe.reg.searchers
    pipe.run([Target("Acme", "acme.com", "Category Manager")])
    assert hits["n"] == 2  # one retry, then skipped
    assert not usage.is_exhausted("tavily")  # rate limits don't burn the month


def test_serper_free_plan_switches_to_plain_queries(usage):
    seen = []

    def serper_strict(kw):
        q = kw["json"]["q"]
        seen.append(q)
        if "site:" in q or '"' in q:
            return Resp(400, {"message": "Query pattern not allowed for free accounts.", "statusCode": 400})
        return Resp(200, {"organic": LINKEDIN_ROWS + [
            {"title": "Acme careers", "link": "https://acme.com/jobs", "snippet": "Category Manager"}]})

    pipe, http = make({"serper.dev": serper_strict, "prospeo.io": prospeo_ok,
                       "reoon.com": reoon_ok}, usage)
    rows = pipe.run([Target("Acme", "acme.com", "Category Manager")])
    assert len(rows) == 2
    assert seen[1] == "Category Manager Acme linkedin"   # plain retry of the first query
    assert all("site:" not in q for q in seen[2:])       # stays plain for the run
    assert usage.used("serper", "lifetime") >= 1         # rejected calls aren't counted
    assert usage.used("serper", "lifetime") == len(seen) - 1


def test_paging_keeps_going_until_the_target(usage):
    def serper_paged(kw):
        page = kw["json"].get("page", 1)
        tag = lambda i: "abcdefghij"[page] + "abcdefghij"[i]   # unique letters per person
        rows = [{"title": f"Amit{tag(i)} Kumar{tag(i)} - Category Manager - Acme | LinkedIn",
                 "link": f"https://in.linkedin.com/in/p{page}-{i}",
                 "snippet": "Category Manager at Acme"} for i in range(5)]
        return Resp(200, {"organic": rows if page <= 4 else []})

    pipe, http = make({"serper.dev": serper_paged, "prospeo.io": prospeo_ok,
                       "reoon.com": lambda kw: Resp(200, {"status": "safe"})}, usage)
    people = pipe.discover(Target("Acme", "acme.com", "Category Manager", limit=12))
    assert len(people) == 12
    assert http.count("serper.dev") == 3  # pages 1-3 of the first query


def test_exports(usage):
    pipe, _ = make({"serper.dev": serper_ok, "prospeo.io": prospeo_ok, "reoon.com": reoon_ok}, usage)
    rows = pipe.run([Target("Acme", "acme.com", "Category Manager")])
    csv_bytes = to_csv(rows)
    assert b"arao@acme.com" in csv_bytes
    xlsx = to_xlsx(rows, pipe.events, pipe.reg.spent_this_run())
    from openpyxl import load_workbook
    import io
    wb = load_workbook(io.BytesIO(xlsx))
    assert wb.sheetnames == ["Results", "Summary", "Run log"]
    assert wb["Results"]["F2"].value == "arao@acme.com"


def test_slow_service_skipped_after_two_timeouts(usage):
    import requests as rq

    def tomba_slow(kw):
        raise rq.ReadTimeout("Read timed out")

    people = [core.Person.from_name(n) for n in ("Ananya Rao", "Kabir Shah", "Ira Sen")]
    pipe, http = make({"prospeo.io": lambda kw: Resp(400, {"error_code": "NO_MATCH"}),
                       "hunter.io": lambda kw: Resp(200, {"data": {}}),
                       "tomba.io": tomba_slow,
                       "reoon.com": lambda kw: Resp(200, {"status": "safe"})}, usage, mode="accurate")
    rows = pipe.run([Target("Acme", "acme.com", people=people)])
    assert len(rows) == 3 and all(r.email for r in rows)   # everyone still gets an answer
    assert http.count("api.tomba.io/v1/email-finder") == 2  # third person skips Tomba
    assert any("timed out twice" in e for e in pipe.events)
    assert usage.used("tomba_search", "month") == 0         # nothing counted


def test_big_targets_use_more_search_variants():
    small = Pipeline.search_variants(Target("Acme", title="Product Manager", limit=20))
    big = Pipeline.search_variants(Target("Acme", title="Product Manager", limit=500,
                                          locations=["Bengaluru", "Mumbai"]))
    assert len(small) == 3
    assert '"Product Manager" "Acme" Bengaluru' in big
    assert '"Senior Product Manager" "Acme"' in big
    assert len(big) == len(set(big)) >= 11


def test_discovery_moves_through_variants_past_googles_page_cap(usage):
    # every search phrasing returns 10 pages of 5 new people, then nothing
    counter = {"n": 0}

    def serper_deep(kw):
        page = kw["json"].get("page", 1)
        if page > 10:
            return Resp(200, {"organic": []})
        rows = []
        for i in range(5):
            counter["n"] += 1
            tag = "".join("abcdefghij"[int(d)] for d in f"{counter['n']:04d}")
            rows.append({"title": f"Amit{tag} Kumar{tag} - Category Manager - Acme | LinkedIn",
                         "link": f"https://in.linkedin.com/in/x{counter['n']}",
                         "snippet": "Category Manager at Acme"})
        return Resp(200, {"organic": rows})

    pipe, http = make({"serper.dev": serper_deep}, usage)
    people = pipe.discover(Target("Acme", "acme.com", "Category Manager", limit=120))
    assert len(people) == 120
    assert http.count("serper.dev") == 24   # 10 pages, 10 pages, then 4 on the third phrasing
