"""CSV and Excel output."""
from __future__ import annotations

import csv
import io
from datetime import datetime

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from .pipeline import CATCH_ALL, FOUND, GUESS, NOT_FOUND, VERIFIED, Row

COLUMNS = ["Company", "Domain", "Name", "Headline", "LinkedIn", "Email", "Status",
           "Confidence", "Found via", "Pattern", "Alternates", "Mail host", "Note"]

STATUS_FILL = {
    VERIFIED: "D9F2E3",
    FOUND: "E3EEFB",
    CATCH_ALL: "FDF0D5",
    GUESS: "EEEEEE",
    NOT_FOUND: "F9DCDC",
}


def to_csv(rows: list[Row]) -> bytes:
    buf = io.StringIO()
    w = csv.DictWriter(buf, fieldnames=COLUMNS)
    w.writeheader()
    for r in rows:
        w.writerow(r.as_dict())
    return buf.getvalue().encode("utf-8-sig")  # BOM so Excel opens it cleanly


def to_xlsx(rows: list[Row], events: list[str] | None = None,
            spent: dict[str, float] | None = None) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Results"
    ws.append(COLUMNS)
    head_font = Font(bold=True, color="FFFFFF")
    head_fill = PatternFill("solid", fgColor="2F3B4C")
    for c in ws[1]:
        c.font, c.fill = head_font, head_fill
        c.alignment = Alignment(vertical="center")

    status_col = COLUMNS.index("Status") + 1
    conf_col = COLUMNS.index("Confidence") + 1
    link_col = COLUMNS.index("LinkedIn") + 1
    for r in rows:
        ws.append([r.as_dict()[c] for c in COLUMNS])
        i = ws.max_row
        fill = STATUS_FILL.get(r.status)
        if fill:
            ws.cell(i, status_col).fill = PatternFill("solid", fgColor=fill)
        ws.cell(i, conf_col).number_format = "0%"
        if r.linkedin:
            cell = ws.cell(i, link_col)
            cell.hyperlink = r.linkedin
            cell.font = Font(color="1F5FBF", underline="single")

    widths = {"Company": 18, "Domain": 18, "Name": 22, "Headline": 38, "LinkedIn": 34,
              "Email": 32, "Status": 18, "Confidence": 11, "Found via": 22, "Pattern": 15,
              "Alternates": 40, "Mail host": 16, "Note": 34}
    for idx, col in enumerate(COLUMNS, start=1):
        ws.column_dimensions[get_column_letter(idx)].width = widths.get(col, 15)
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions

    summary = wb.create_sheet("Summary")
    summary.append(["Generated", datetime.now().strftime("%Y-%m-%d %H:%M")])
    summary.append([])
    summary.append(["Status", "Count"])
    for s in (VERIFIED, FOUND, CATCH_ALL, GUESS, NOT_FOUND):
        summary.append([s, sum(1 for r in rows if r.status == s)])
    summary.append([])
    summary.append(["Credits used this run", ""])
    for name, n in (spent or {}).items():
        summary.append([name, n])
    for row in summary.iter_rows():
        if row[0].value in ("Status", "Credits used this run"):
            row[0].font = Font(bold=True)
    summary.column_dimensions["A"].width = 28
    summary.column_dimensions["B"].width = 20

    if events:
        log = wb.create_sheet("Run log")
        for line in events:
            log.append([line])
        log.column_dimensions["A"].width = 120

    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()
